#!/usr/bin/env python3
"""Export worksheets to position-preserving Markdown. Python 3.10+."""
import argparse
import datetime as dt
from fractions import Fraction
import html
import json
from pathlib import Path
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def safe_name(name):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name).strip(' .')
    if not name:
        name = 'sheet'
    if name.split('.')[0].upper() in {'CON', 'PRN', 'AUX', 'NUL',
            *(f'COM{i}' for i in range(1, 10)), *(f'LPT{i}' for i in range(1, 10))}:
        name = '_' + name
    return name[:100]


def text_value(cell):
    """Handle common recipe quantity formats; leave unsupported formats raw."""
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, (int, float)):
        fmt = cell.number_format.split(';')[0]
        if re.fullmatch(r'# ?\?+/\?+', fmt):
            denominator = 10 ** len(fmt.split('/')[1]) - 1
            fraction = Fraction(abs(value)).limit_denominator(denominator)
            whole, remainder = divmod(fraction.numerator, fraction.denominator)
            result = str(whole) if not remainder else (
                (str(whole) + ' ' if whole else '') +
                f'{remainder}/{fraction.denominator}')
            return ('-' if value < 0 else '') + result
        if re.fullmatch(r'0(?:\.0+)?%', fmt):
            places = len(fmt.split('.')[1].rstrip('%')) if '.' in fmt else 0
            return f'{value * 100:.{places}f}%'
        if re.fullmatch(r'0+', fmt):
            return f'{value:0{len(fmt)}.0f}'
        if re.fullmatch(r'0\.0+', fmt):
            return f'{value:.{len(fmt.split(".")[1])}f}'
        return str(value)
    return str(value)


def escaped(text):
    return html.escape(text, quote=True).replace('\r\n', '\n').replace('\r', '\n').replace('\n', '<br>')


def markdown_cell(text):
    # Entities prevent cell contents from becoming Markdown or breaking columns.
    text = escaped(text)
    for char in '\\|`*_[]~':
        text = text.replace(char, f'&#{ord(char)};')
    return text


def render_sheet(sheet, cached_sheet, mode='html', formulas='formula', max_cells=250000):
    cells = {}
    warnings = []
    # Ignore style-only trailing cells while retaining interior/leading gaps.
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                text = text_value(cell)
                if cell.data_type == 'f' and formulas == 'cached':
                    cached = cached_sheet[cell.coordinate]
                    if cached.value is None:
                        warnings.append(f'{cell.coordinate}: no cached result; preserved formula')
                    else:
                        text = text_value(cached)
                cells[(cell.row, cell.column)] = text
    merges = list(sheet.merged_cells.ranges)
    rows = max((r for r, c in cells), default=0)
    cols = max((c for r, c in cells), default=0)
    if cells:
        for merged in merges:
            if (merged.min_row, merged.min_col) in cells:
                rows = max(rows, merged.max_row)
                cols = max(cols, merged.max_col)
    if rows * cols > max_cells:
        raise ValueError(f'{sheet.title}: grid has {rows * cols:,} cells; limit is {max_cells:,}')
    title = markdown_cell(sheet.title)
    lines = [f'# {title}', '', f'<!-- Worksheet state: {sheet.sheet_state}; grid starts at A1. -->', '']
    if not cells:
        lines += ['_Empty worksheet._', '']
        return '\n'.join(lines), warnings
    if mode == 'markdown':
        lines.append('| ' + ' | '.join(get_column_letter(c) for c in range(1, cols + 1)) + ' |')
        lines.append('| ' + ' | '.join('---' for _ in range(cols)) + ' |')
        for r in range(1, rows + 1):
            lines.append('| ' + ' | '.join(markdown_cell(cells.get((r, c), '')) for c in range(1, cols + 1)) + ' |')
        if merges:
            lines += ['', '<!-- Merged ranges: ' + ', '.join(str(m) for m in merges) + '; values remain in the top-left cell. -->']
    else:
        anchors, covered = {}, set()
        for merged in merges:
            if merged.min_row > rows or merged.min_col > cols:
                continue
            anchors[(merged.min_row, merged.min_col)] = merged
            for r in range(merged.min_row, min(rows, merged.max_row) + 1):
                for c in range(merged.min_col, min(cols, merged.max_col) + 1):
                    if (r, c) != (merged.min_row, merged.min_col):
                        covered.add((r, c))
        lines.append('<table>')
        for r in range(1, rows + 1):
            lines.append('  <tr>')
            for c in range(1, cols + 1):
                if (r, c) in covered:
                    continue
                merged = anchors.get((r, c))
                attrs = ''
                if merged:
                    attrs = f' rowspan="{min(rows, merged.max_row) - r + 1}" colspan="{min(cols, merged.max_col) - c + 1}"'
                lines.append(f'    <td{attrs}>{escaped(cells.get((r, c), ""))}</td>')
            lines.append('  </tr>')
        lines.append('</table>')
    lines.append('')
    if sheet._images or sheet._charts:
        warnings.append('Embedded images/charts are not exported; keep the original workbook.')
    return '\n'.join(lines), warnings


def convert(source, output, mode='html', formulas='formula', overwrite=False, max_cells=250000):
    source, output = Path(source), Path(output)
    if source.suffix.lower() not in {'.xlsx', '.xlsm'}:
        raise ValueError('Use an .xlsx or .xlsm workbook; save legacy .xls as .xlsx first.')
    workbook = load_workbook(source, data_only=False)
    cached = load_workbook(source, data_only=True) if formulas == 'cached' else None
    try:
        names = set()
        exports = []
        for index, sheet in enumerate(workbook.worksheets, 1):
            # Prefix preserves workbook order and avoids duplicate sanitized names.
            name = f'{index:03d}-{safe_name(sheet.title)}.md'
            if name.casefold() in names:
                raise ValueError(f'Duplicate output filename: {name}')
            names.add(name.casefold())
            if sheet.max_row * sheet.max_column > max_cells:
                raise ValueError(f'{sheet.title}: worksheet dimensions exceed --max-cells; inspect stray formatting or raise the limit.')
            content, warnings = render_sheet(sheet, cached[sheet.title] if cached else None,
                                              mode, formulas, max_cells)
            exports.append((name, content, sheet.title, warnings))
        targets = [output / item[0] for item in exports] + [output / 'conversion-report.json']
        if not overwrite:
            existing = [str(p) for p in targets if p.exists()]
            if existing:
                raise FileExistsError('Refusing to overwrite: ' + ', '.join(existing))
        output.mkdir(parents=True, exist_ok=True)
        report = {'source': source.name, 'mode': mode, 'formulas': formulas, 'sheets': []}
        for name, content, title, warnings in exports:
            (output / name).write_text(content, encoding='utf-8')
            report['sheets'].append({'sheet': title, 'file': name, 'warnings': warnings})
        (output / 'conversion-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        return report
    finally:
        workbook.close()
        if cached:
            cached.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('-o', '--output', type=Path, default=Path('markdown'))
    parser.add_argument('--mode', choices=['html', 'markdown'], default='html')
    parser.add_argument('--formulas', choices=['formula', 'cached'], default='formula')
    parser.add_argument('--overwrite', action='store_true')
    parser.add_argument('--max-cells', type=int, default=250000)
    args = parser.parse_args()
    try:
        if args.max_cells < 1:
            raise ValueError('--max-cells must be positive')
        report = convert(args.workbook, args.output, args.mode, args.formulas, args.overwrite, args.max_cells)
    except Exception as exc:
        print(f'Error: {exc}', file=sys.stderr)
        return 1
    for sheet in report['sheets']:
        print(sheet['file'])
        for warning in sheet['warnings']:
            print(f"Warning: {sheet['sheet']}: {warning}", file=sys.stderr)
    print(f"Exported {len(report['sheets'])} worksheets to {args.output}")
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
